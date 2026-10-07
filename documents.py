"""Document handling for chat: text cleanup, chunking, outline, retrieval.

Deliberately free of heavy imports (no llama_cpp, no FastAPI) so it can be unit
tested without the model. Anything that needs a tokenizer takes a
`count_tokens` callable; app.py wires in the real one.

How it fits together:

- The CLIENT owns the documents (same pattern as the compaction summary). On
  upload or paste the server only extracts, cleans and chunks the text
  (`build_document`) and returns the result; the browser stores it and sends
  the chat's documents with every chat request. The server stays stateless, so
  any Cloud Run instance can answer, restarts lose nothing.
- Per question, `select_context` picks what goes into the prompt, within a
  token budget: an outline plus samples for overview questions ("summarize
  this"), the named section for "section 12", otherwise BM25 over the chunks.
"""

from __future__ import annotations

import bisect
import hashlib
import io
import math
import re
import unicodedata
from collections import Counter
from typing import Callable

CountTokens = Callable[[str], int]

CHUNK_TARGET = 900  # characters per chunk (before the small overlap prefix)
CHUNK_OVERLAP = 100
SAMPLE_CHARS = 320  # one overview excerpt
MAX_DOCS = 20
MAX_NAME_CHARS = 200
MAX_CHUNK_CHARS = 3000  # validation cap for one chunk sent by a client
MAX_DOC_CHARS = 400_000  # per chat: about 80 pages; the model can't use more than a few thousand tokens anyway
MAX_TOTAL_CHARS = 420_000  # validation cap for a request (a little above MAX_DOC_CHARS)
MAX_TOTAL_CHUNKS = 1500
MAX_OUTLINE = 300
MIN_BUDGET_TOKENS = 120  # below this there is no room for document text, only a note
CHARS_PER_TOKEN = 4  # only used to size things before the real tokenizer measures them
CHARS_PER_PAGE = 5000  # only for the page counts quoted in messages (400,000 characters is "about 80 pages")


class DocumentError(Exception):
    """A document problem the caller turns into an HTTP error (status + message)."""

    def __init__(self, status: int, message: str):
        super().__init__(message)
        self.status = status
        self.message = message


# ---------------------------------------------------------------------------
# Cleaning
# ---------------------------------------------------------------------------

_CONTROL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f-\x9f]")
_INVISIBLE = re.compile("[\u200b-\u200d\u2060\ufeff\u00ad]")  # zero-width characters and soft hyphens
_PAGE_NUMBER = re.compile(r"^\s*(?:page\s+)?\d{1,4}(?:\s*(?:/|of)\s*\d{1,4})?\s*$", re.I)
_BULLET = re.compile(r"^\s*(?:[-*\u2022\u25cf\u25aa\u25e6\u2023\u2013\u2014]\s+|\(?\d{1,3}[.)]\s+|\(?[A-Za-z][.)]\s+)")
_NUM_HEADING = re.compile(
    r"^(?:(?:Section|SECTION|Part|PART|Chapter|CHAPTER|Article|ARTICLE)\s+)?"
    r"\d{1,3}(?:\.\d{1,3}){0,3}[.):]?\s+[A-Z][^\n]{1,90}$"
)


def _is_heading(line: str) -> bool:
    """A numbered heading ("5. BLUE AI", "3.2 Verification") or a short
    ALL-CAPS line. Sentences end in punctuation and are longer, so they don't
    count."""
    s = line.strip()
    if not 3 <= len(s) <= 100 or s.endswith((".", ",", ";", "!", "?")):
        return False
    if _NUM_HEADING.match(s):
        return True
    letters = [c for c in s if c.isalpha()]
    return (
        len(s) <= 80
        and len(letters) >= 3
        and sum(c.isupper() for c in letters) / len(letters) >= 0.85
        and len(s.split()) <= 12
    )


def _rejoin_paragraph(lines: list[str]) -> str:
    """Undo PDF hard wrapping inside one paragraph: join wrapped lines (and
    de-hyphenate "veri-\\nfication"), but keep headings and list items on their
    own lines."""
    out: list[str] = []
    cur = ""
    for raw in lines:
        line = raw.strip()
        if not line or _PAGE_NUMBER.match(line):
            continue
        if not cur:
            cur = line
        elif _is_heading(cur) or _is_heading(line) or _BULLET.match(line):
            out.append(cur)
            cur = line
        elif len(cur) > 1 and cur.endswith("-") and cur[-2].isalpha() and line[:1].islower():
            cur = cur[:-1] + line
        else:
            cur = f"{cur} {line}"
    if cur:
        out.append(cur)
    return "\n".join(out)


def clean_text(raw: str) -> str:
    """Normalize extracted text so matching and chunking behave.

    NFKC turns ligatures into letters ("\ufb01nancial" -> "financial", which a
    keyword tokenizer would otherwise split in two), then control and
    zero-width characters go, wrapped lines are rejoined, and blank-line
    paragraphs are kept."""
    text = unicodedata.normalize("NFKC", raw or "")
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace("\t", " ").replace("\u00a0", " ")
    text = _INVISIBLE.sub("", text)
    text = _CONTROL.sub("", text)
    text = re.sub(r"[ \f\v]+", " ", text)
    paragraphs = (_rejoin_paragraph(p.split("\n")) for p in re.split(r"\n\s*\n", text))
    return "\n\n".join(p for p in paragraphs if p).strip()


# ---------------------------------------------------------------------------
# Chunking and outline
# ---------------------------------------------------------------------------

_CUT_ORDER = ("\n", ". ", "! ", "? ", "; ", ": ", ", ", " ")


def _split_long(text: str, s: int, e: int, target: int) -> list[tuple[int, int]]:
    """Cut text[s:e] into spans of at most `target` characters, at a line end,
    sentence end, clause or word boundary, never mid-word (unless one single
    token is longer than the target, like a URL)."""
    spans: list[tuple[int, int]] = []
    while e - s > target:
        window_end = s + target
        lo = s + target // 2  # prefer cuts in the back half of the window
        end = -1
        for pat in _CUT_ORDER:
            idx = text.rfind(pat, lo, window_end)
            if idx != -1:
                end = idx + (0 if pat in ("\n", " ") else 1)  # keep the punctuation, drop the space
                break
        if end <= s:
            end = window_end
        spans.append((s, end))
        s = end
        while s < e and text[s].isspace():
            s += 1
    if e > s:
        spans.append((s, e))
    return spans


def _spans(text: str, target: int) -> list[tuple[int, int]]:
    spans: list[tuple[int, int]] = []
    pos = 0
    pieces: list[tuple[int, int]] = []
    for m in re.finditer(r"\n\s*\n", text):
        pieces.append((pos, m.start()))
        pos = m.end()
    pieces.append((pos, len(text)))
    for s, e in pieces:
        while s < e and text[s].isspace():
            s += 1
        while e > s and text[e - 1].isspace():
            e -= 1
        if e > s:
            spans.extend(_split_long(text, s, e, target))
    return spans


def chunk_text(text: str, target: int = CHUNK_TARGET, overlap: int = CHUNK_OVERLAP) -> list[dict]:
    """Chunks of about `target` characters on paragraph/sentence/word
    boundaries. Each is {"text", "ov", "start", "end"}: `text` begins with an
    `ov`-character overlap prefix (the tail of the previous chunk, so a sentence
    split across a boundary still matches); start/end are the offsets of the
    chunk's own content in `text`."""
    groups: list[tuple[int, int]] = []
    first = last = None
    for s, e in _spans(text, target):
        if first is not None and e - first > target:
            groups.append((first, last))
            first = None
        if first is None:
            first = s
        last = e
    if first is not None:
        groups.append((first, last))

    out: list[dict] = []
    for k, (s, e) in enumerate(groups):
        core = text[s:e]
        prefix = ""
        if k > 0 and overlap > 0:
            prev = text[groups[k - 1][0] : groups[k - 1][1]]
            tail = prev[-overlap:]
            if len(prev) > overlap:
                m = re.search(r"\s", tail)  # start at a word boundary
                tail = tail[m.end() :] if m else ""
            prefix = tail.strip()
        out.append(
            {
                "text": f"{prefix} {core}" if prefix else core,
                "ov": len(prefix) + 1 if prefix else 0,
                "start": s,
                "end": e,
            }
        )
    return out


def build_outline(text: str, chunks: list[dict]) -> list[dict]:
    """Headings mapped to the chunk that contains them: [{"t": title, "c": chunk_index}].
    A heading repeated 3+ times without a number is a running page header, not
    a section, and is dropped."""
    if not chunks:
        return []
    found: list[tuple[int, str]] = []
    pos = 0
    for line in text.split("\n"):
        if _is_heading(line):
            found.append((pos, line.strip()))
        pos += len(line) + 1
    repeats = Counter(t for _, t in found)
    starts = [c["start"] for c in chunks]
    ends = [c["end"] for c in chunks]
    outline: list[dict] = []
    for offset, title in found:
        if repeats[title] >= 3 and not _NUM_HEADING.match(title):
            continue
        idx = max(0, bisect.bisect_right(starts, offset) - 1)
        if offset >= ends[idx] and idx + 1 < len(chunks):  # heading sits in the gap before the next chunk
            idx += 1
        outline.append({"t": title[:120], "c": idx})
        if len(outline) >= MAX_OUTLINE:
            break
    return outline


def doc_id_for(clean: str) -> str:
    return hashlib.sha1(clean.encode("utf-8")).hexdigest()[:16]


def clean_name(name: str | None, fallback: str = "Document") -> str:
    cleaned = " ".join((name or "").split())[:MAX_NAME_CHARS]
    return cleaned or fallback


def build_document(name: str | None, raw_text: str, *, pages: int | None = None) -> dict:
    """Everything the client stores for one document:
    {doc_id, name, chars, chunks[], overlaps[], outline[], warnings[]}.
    The id is a hash of the cleaned text, so the same file is the same document."""
    text = clean_text(raw_text)
    if not text:
        if pages:
            raise DocumentError(
                400,
                "This PDF has no selectable text (it looks scanned or image-only). OCR isn't supported, "
                "so I can't read it.",
            )
        raise DocumentError(400, "No text found in that file.")
    if len(text) > MAX_DOC_CHARS:
        raise DocumentError(
            413,
            f"That document is too large (about {len(text) // CHARS_PER_PAGE} pages of text; the limit is about "
            f"{MAX_DOC_CHARS // CHARS_PER_PAGE} pages per chat). Upload a shorter file or split it.",
        )
    chunks = chunk_text(text)
    warnings: list[str] = []
    if pages and len(text) / pages < 200:
        warnings.append(
            f"Very little text for {pages} page(s) ({len(text) // pages} characters per page). "
            "Parts of this PDF may be scanned images, which I can't read."
        )
    return {
        "doc_id": doc_id_for(text),
        "name": clean_name(name),
        "chars": len(text),
        "chunks": [c["text"] for c in chunks],
        "overlaps": [c["ov"] for c in chunks],
        "outline": build_outline(text, chunks),
        "warnings": warnings,
    }


# ---------------------------------------------------------------------------
# Extraction from files
# ---------------------------------------------------------------------------


def extract_file(filename: str, data: bytes) -> tuple[str, dict]:
    """(raw text, {"pages": n} for PDFs). Raises ValueError for an unsupported
    type; parser errors propagate (the caller turns both into HTTP 400)."""
    name = filename.lower()
    if name.endswith(".pdf"):
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            reader.decrypt("")
        texts = [page.extract_text() or "" for page in reader.pages]
        return "\n\n".join(texts), {"pages": len(texts)}
    if name.endswith(".docx"):
        from docx import Document

        doc = Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text.strip() for cell in row.cells))
        return "\n\n".join(parts), {}
    if name.endswith((".txt", ".md")):
        return data.decode("utf-8", errors="replace"), {}
    raise ValueError("Unsupported file type. Use PDF, DOCX, TXT or MD.")


# ---------------------------------------------------------------------------
# Validating what a client sends with a chat request
# ---------------------------------------------------------------------------


def validate_documents(raw: object) -> list[dict]:
    """The documents of a chat request, checked and normalized. DocumentError
    (400 for a bad shape, 413 for too much) otherwise."""
    if raw is None:
        return []
    if not isinstance(raw, list):
        raise DocumentError(400, "documents must be a list.")
    if len(raw) > MAX_DOCS:
        raise DocumentError(413, f"Too many documents attached (the limit is {MAX_DOCS}). Remove some.")
    docs: list[dict] = []
    seen: set[str] = set()
    total_chars = total_chunks = 0
    for d in raw:
        if not isinstance(d, dict):
            raise DocumentError(400, "Each document must be an object.")
        doc_id, name, chunks = d.get("id"), d.get("name"), d.get("chunks")
        if not isinstance(doc_id, str) or not doc_id or len(doc_id) > 64:
            raise DocumentError(400, "A document has no valid id.")
        if not isinstance(chunks, list) or not chunks or not all(isinstance(c, str) for c in chunks):
            raise DocumentError(400, "A document has no valid chunks.")
        if doc_id in seen:
            continue
        seen.add(doc_id)
        if any(len(c) > MAX_CHUNK_CHARS for c in chunks):
            raise DocumentError(413, "A document chunk is too large.")
        overlaps = d.get("overlaps")
        if not (isinstance(overlaps, list) and len(overlaps) == len(chunks)):
            overlaps = [0] * len(chunks)
        overlaps = [o if isinstance(o, int) and 0 <= o < len(chunks[i]) else 0 for i, o in enumerate(overlaps)]
        total_chunks += len(chunks)
        # The document's own text, not each chunk's duplicated overlap prefix: this is the number the
        # client limits (cleaned characters), so a chat the UI accepts is never refused here.
        total_chars += sum(len(c) - o for c, o in zip(chunks, overlaps))
        if total_chars > MAX_TOTAL_CHARS or total_chunks > MAX_TOTAL_CHUNKS:
            raise DocumentError(
                413,
                "The attached documents are too large together (the limit is about "
                f"{MAX_DOC_CHARS // CHARS_PER_PAGE} pages per chat). Remove one and try again.",
            )
        outline = []
        for item in d.get("outline") or []:
            if isinstance(item, dict) and isinstance(item.get("t"), str) and isinstance(item.get("c"), int):
                if 0 <= item["c"] < len(chunks):
                    outline.append({"t": item["t"][:120], "c": item["c"]})
        docs.append(
            {
                "id": doc_id,
                "name": clean_name(name if isinstance(name, str) else None),
                "chunks": chunks,
                "overlaps": overlaps,
                "outline": outline[:MAX_OUTLINE],
            }
        )
    return docs


# ---------------------------------------------------------------------------
# Retrieval
# ---------------------------------------------------------------------------

_WORD = re.compile(r"[^\W_]+", re.UNICODE)

STOPWORDS = frozenset(
    "a an the and or but if of to in on at by for with about as is are was were be been being this that these those "
    "it its from into than then so such not no nor do does did done have has had having i you he she we they me him "
    "her us them my your his our their what which who whom whose when where why how can could should would may might "
    "will shall also just very more most some any each other over under up out off again further once there here all "
    "both few many much own same too only s t".split()
)
_CHITCHAT = frozenset(
    "hi hello hey hiya thanks thank thx ok okay oka bye goodbye yes yeah yep nope please cool great good nice morning "
    "evening afternoon night lol sure alright".split()
)


def _stem(w: str) -> str:
    """Deliberately light: enough for plural/verb endings ("services" and
    "service", "verifies" and "verify"), not a real stemmer."""
    if len(w) > 4 and w.endswith("ies"):
        w = w[:-3] + "y"
    else:
        for suf in ("ing", "ed", "es", "s"):
            if w.endswith(suf) and len(w) - len(suf) >= 3 and not (suf == "s" and w.endswith("ss")):
                w = w[: -len(suf)]
                break
    if len(w) > 3 and w.endswith("e"):
        w = w[:-1]
    return w


def _terms(text: str) -> list[str]:
    tokens = _WORD.findall(unicodedata.normalize("NFKC", text).lower())
    return [_stem(t) for t in tokens if t not in STOPWORDS and (len(t) > 1 or t.isdigit())]


_CHITCHAT_STEMS = frozenset(_stem(w) for w in _CHITCHAT)

_OVERVIEW = re.compile(
    r"\b(?:summari[sz]e|summari[sz]ation|summary|overview|tl;?dr|gist|in a nutshell|table of contents|"
    r"(?:main|key|major) (?:points?|ideas?|takeaways?|topics?|themes?)|"
    r"what(?:'s|\s+is|\s+are)\s+(?:this|these|that|the)\b[^?.!]{0,40}\b(?:about|say|cover|contain)|"
    r"what(?:'s|\s+is)\s+in\s+(?:this|the|my)\b|"
    r"what\s+(?:does|do)\s+(?:this|these|that|the)\b[^?.!]{0,40}\b(?:say|cover|contain|talk about|discuss)|"
    r"(?:explain|describe|walk me through|go over|read|review|tell me about)\s+(?:this|the|my|that|these)\b"
    r"[^?.!]{0,30}\b(?:document|doc|pdf|file|paper|text|article|attachment|report|paste|pasted|content))\b",
    re.I,
)
_SECTION_REF = re.compile(r"\b(?:sections?|parts?|chapters?|clauses?|articles?)\s*#?\s*(\d{1,3}(?:\.\d{1,3}){0,3})\b", re.I)
_FOLLOWUP_MAX_TERMS = 2


def is_overview_question(question: str) -> bool:
    return bool(_OVERVIEW.search(question))


def _section_hit(documents: list[dict], question: str) -> tuple[int, int, int] | None:
    """(doc index, first chunk, one past the last chunk) of the section the
    question names ("section 12", "part 3.2"), found through the outline."""
    m = _SECTION_REF.search(question)
    if not m:
        return None
    num = re.escape(m.group(1))
    pat = re.compile(
        rf"^(?:(?:section|part|chapter|article)\s+)?{num}(?!\d)(?!\.\d)[.):\-\s]", re.I
    )
    for di, doc in enumerate(documents):
        outline = doc["outline"]
        for oi, item in enumerate(outline):
            if pat.match(item["t"]):
                start = item["c"]
                nxt = next((o["c"] for o in outline[oi + 1 :] if o["c"] > start), len(doc["chunks"]))
                return di, start, max(start + 1, nxt)
    return None


def _rank(documents: list[dict], terms: list[str]) -> list[tuple[float, int, int]]:
    """Okapi BM25 over every chunk of every document; (score, doc, chunk),
    best first, only chunks that match at least one query term."""
    if not terms:
        return []
    entries: list[tuple[int, int, Counter, int]] = []
    df: Counter = Counter()
    for di, doc in enumerate(documents):
        for ci, chunk in enumerate(doc["chunks"]):
            toks = _terms(chunk)
            tf = Counter(toks)
            entries.append((di, ci, tf, len(toks)))
            df.update(tf.keys())
    n = len(entries)
    if n == 0:
        return []
    avgdl = sum(e[3] for e in entries) / n or 1.0
    k1, b = 1.5, 0.75
    ranked: list[tuple[float, int, int]] = []
    for di, ci, tf, length in entries:
        score = 0.0
        for term in set(terms):
            f = tf.get(term, 0)
            if not f:
                continue
            idf = math.log((n - df[term] + 0.5) / (df[term] + 0.5) + 1)
            score += idf * (f * (k1 + 1)) / (f + k1 * (1 - b + b * length / avgdl))
        if score > 0:
            ranked.append((score, di, ci))
    ranked.sort(key=lambda r: (-r[0], r[1], r[2]))
    return ranked


def _core(doc: dict, ci: int) -> str:
    """Chunk text without its overlap prefix."""
    return doc["chunks"][ci][doc["overlaps"][ci] :]


def _join(doc: dict, indices: list[int]) -> str:
    """Selected chunks of one document in reading order, adjacent chunks merged
    without their duplicated overlap, gaps marked."""
    parts: list[str] = []
    prev = None
    for ci in sorted(set(indices)):
        if prev is not None and ci == prev + 1:
            parts.append(" " + _core(doc, ci))
        elif prev is None:
            parts.append(doc["chunks"][ci] if ci == 0 else _core(doc, ci))
        else:
            parts.append("\n[…]\n" + _core(doc, ci))
        prev = ci
    return "".join(parts)


def _clip(text: str, n: int) -> str:
    if len(text) <= n:
        return text
    cut = text.rfind(" ", 0, n)
    return text[: cut if cut > n * 0.6 else n].rstrip() + " […]"


def _trim_to_tokens(text: str, max_tokens: int, count: CountTokens) -> str:
    for _ in range(8):
        used = count(text)
        if used <= max_tokens:
            return text
        text = _clip(text, max(60, int(len(text) * max_tokens / used * 0.93)))
    return text


def _even_indices(n: int, k: int) -> list[int]:
    if k >= n:
        return list(range(n))
    if k <= 1:
        return [0]
    return sorted({round(i * (n - 1) / (k - 1)) for i in range(k)})


def _outline_text(outline: list[dict], max_chars: int) -> str:
    if not outline or max_chars < 40:
        return ""
    text = "Outline: " + "; ".join(item["t"] for item in outline)
    return text if len(text) <= max_chars else _clip(text, max_chars)


def _overview_body(doc: dict, share_chars: int) -> str:
    outline = _outline_text(doc["outline"], int(share_chars * 0.25))
    avail = max(200, share_chars - len(outline))
    idxs = _even_indices(len(doc["chunks"]), max(3, avail // (SAMPLE_CHARS + 10)))
    per_sample = max(100, min(CHUNK_TARGET, avail // len(idxs)))
    samples = [_clip(_core(doc, i), per_sample) for i in idxs]
    return (outline + "\n" if outline else "") + "\n[…]\n".join(samples)


def _note_for(names: list[str], what: str) -> str:
    return f"The attached documents ({', '.join(names)}) were searched and {what}"


def select_context(
    documents: list[dict],
    question: str,
    *,
    prev_question: str = "",
    budget_tokens: int,
    count_tokens: CountTokens | None = None,
) -> dict:
    """What to put in the prompt for this question.

    Returns {"text", "mode", "selected", "total", "doc_ids"} where mode is one
    of: none (greeting-like, nothing added), overview, section, search,
    followup, nomatch, nobudget. `documents` must come from validate_documents.
    """
    count = count_tokens or (lambda t: max(1, len(t) // CHARS_PER_TOKEN))
    total = sum(len(d["chunks"]) for d in documents)
    info = {"text": "", "mode": "none", "selected": 0, "total": total, "doc_ids": [d["id"] for d in documents]}
    if not documents:
        return info
    names = [d["name"] for d in documents]
    question = question.strip()

    section = _section_hit(documents, question)
    overview = not section and is_overview_question(question)
    q_terms = _terms(question)
    content = [t for t in q_terms if t not in _CHITCHAT_STEMS]
    if not section and not overview and not content:
        return info  # "Hi", "thanks": the documents are not part of this turn

    if budget_tokens < MIN_BUDGET_TOKENS:
        info.update(mode="nobudget", text=_note_for(names, "there was no room to include them in this prompt. Say so rather than guessing."))
        return info

    if section:
        di, start, end = section
        doc = documents[di]
        picked = list(range(start, min(end, start + 3)))
        body = f"[{doc['name']}]\n{_join(doc, picked)}"
        info.update(mode="section", selected=len(picked), doc_ids=[doc["id"]])
        info["text"] = _trim_to_tokens(body, budget_tokens, count)
        return info

    if overview:
        note = (
            "This is the outline plus excerpts sampled across the whole document(s), not the full text. "
            "Base the answer on them and say it is based on excerpts."
        )
        note_tokens = count(note)
        share_chars = int(max(200, budget_tokens - note_tokens) * CHARS_PER_TOKEN * 0.9 / len(documents))
        text = ""
        for _ in range(5):
            blocks = [f"[{d['name']}]\n{_overview_body(d, share_chars)}" for d in documents]
            text = note + "\n\n" + "\n\n---\n\n".join(blocks)
            if count(text) <= budget_tokens:
                break
            share_chars = int(share_chars * 0.82)
        info.update(
            mode="overview",
            text=_trim_to_tokens(text, budget_tokens, count),
            selected=sum(min(len(d["chunks"]), max(3, share_chars // (SAMPLE_CHARS + 10))) for d in documents),
        )
        return info

    terms, mode = content, "search"
    if len(content) <= _FOLLOWUP_MAX_TERMS and prev_question:
        terms = content + [t for t in _terms(prev_question) if t not in _CHITCHAT_STEMS]
        mode = "followup"
    ranked = _rank(documents, terms)
    if not ranked:
        info.update(mode="nomatch", text=_note_for(names, "no passage matched this question. Say you couldn't find it in the documents instead of guessing."))
        return info

    order: list[tuple[int, int]] = []
    _, best_doc, best_chunk = ranked[0]
    order.append((best_doc, best_chunk))
    neighbour = best_chunk + 1 if best_chunk + 1 < len(documents[best_doc]["chunks"]) else best_chunk - 1
    if neighbour >= 0 and neighbour != best_chunk:
        order.append((best_doc, neighbour))
    for _, di, ci in ranked[1:]:
        if len(order) >= 5:
            break
        if (di, ci) not in order:
            order.append((di, ci))

    def render(pairs: list[tuple[int, int]]) -> str:
        by_doc: dict[int, list[int]] = {}
        for di, ci in pairs:
            by_doc.setdefault(di, []).append(ci)
        return "\n\n---\n\n".join(f"[{documents[di]['name']}]\n{_join(documents[di], idx)}" for di, idx in sorted(by_doc.items()))

    keep = len(order)
    while keep > 1 and count(render(order[:keep])) > budget_tokens:
        keep -= 1
    text = _trim_to_tokens(render(order[:keep]), budget_tokens, count)
    info.update(mode=mode, text=text, selected=keep, doc_ids=sorted({documents[di]["id"] for di, _ in order[:keep]}))
    return info
